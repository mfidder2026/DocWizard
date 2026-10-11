"""python-pptx backend: SlideScenes -> a native, fully editable .pptx.

Real text boxes (exact line pitch, no autofit, zero insets so measurement == rendering), native tables and
charts styled with the theme, speaker notes, slide titles in title placeholders (outline view and screen
readers), the theme's colours/fonts written into the file's own theme so new slides match.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .scene import ChartP, Line, Para, Picture, Shape, SlideScene, TableP, TextBox, emu
from .spec import DeckSpec
from .themes import A, Theme

if TYPE_CHECKING:  # pragma: no cover
    from .template import TemplateInfo

SLIDE_W, SLIDE_H = 12192000, 6858000  # 13.333 x 7.5 in
_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
_SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
           "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON}
_PH_INSET_X, _PH_INSET_Y = 91440, 45720  # default bodyPr insets (7.2 pt / 3.6 pt)
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _rgb(c: str) -> RGBColor:
    return RGBColor.from_string(c.lstrip("#").upper())


def _strip_style(shape) -> None:
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


# --------------------------------------------------------------------------- theme part


def apply_theme_part(prs, theme: Theme) -> None:
    """Write the theme's colours and fonts into the presentation theme (so new slides/charts match)."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    root.set("name", f"DocWizard {theme.label}")
    scheme = root.find(f".//{{{A}}}clrScheme")
    if scheme is not None:
        scheme.set("name", f"DocWizard {theme.label}")
        bg, txt = theme.background, theme.text
        pal = list(theme.chart) + [theme.accent] * 6
        values = {"dk1": bg if theme.dark else txt, "lt1": txt if theme.dark else bg, "dk2": theme.accent,
                  "lt2": theme.surface, "hlink": theme.accent, "folHlink": theme.muted}
        for i in range(6):
            values[f"accent{i + 1}"] = pal[i]
        for el in scheme:
            name = etree.QName(el).localname
            if name in values:
                for child in list(el):
                    el.remove(child)
                etree.SubElement(el, f"{{{A}}}srgbClr", val=values[name].lstrip("#").upper())
    for tag, face in (("majorFont", theme.heading_font), ("minorFont", theme.body_font)):
        latin = root.find(f".//{{{A}}}{tag}/{{{A}}}latin")
        if latin is not None:
            latin.set("typeface", face)
            for k in ("panose", "pitchFamily", "charset"):
                latin.attrib.pop(k, None)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if theme.dark:
        cm = prs.slide_master._element.find(qn("p:clrMap"))
        if cm is not None:
            cm.set("bg1", "dk1")
            cm.set("tx1", "lt1")
            cm.set("bg2", "dk2")
            cm.set("tx2", "lt2")


# --------------------------------------------------------------------------- text


def _body_pr(tf, anchor: str, zero_insets: bool) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if zero_insets:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR.get(anchor, MSO_ANCHOR.TOP)


def _insert_bullet_xml(pPr, para: Para, inherit: bool) -> None:
    # bullet elements must precede tabLst/defRPr/extLst
    succ = None
    for tag in ("a:tabLst", "a:defRPr", "a:extLst"):
        succ = pPr.find(qn(tag))
        if succ is not None:
            break

    def add(el):
        if succ is not None:
            succ.addprevious(el)
        else:
            pPr.append(el)

    if para.bullet == "none" or (not para.bullet and not inherit):
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        add(etree.Element(qn("a:buNone")))
        return
    if not para.bullet:
        return
    pPr.set("marL", str(emu(para.indent)))
    pPr.set("indent", str(-emu(para.indent - para.bullet_at)))
    if para.bullet == "inherit":
        return
    clr = etree.Element(qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr"), val=para.bullet_color.lstrip("#").upper())
    add(clr)
    add(etree.Element(qn("a:buSzPct"), val="70000" if para.level == 0 else "100000"))
    add(etree.Element(qn("a:buFont"), typeface="Arial"))
    add(etree.Element(qn("a:buChar"), char=para.bullet))


def _run(p, text: str, para: Para, bold: bool, lang: str, inherit: bool) -> None:
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(para.size)
    if not inherit:
        f.name = para.family
        f.color.rgb = _rgb(para.color)
        f.bold = bold
        f.italic = para.italic
    else:
        if bold:
            f.bold = True
        if para.bullet == "none" and para.bold:  # column headings inside template placeholders
            f.color.rgb = _rgb(para.color)
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", lang)
    if para.spacing:
        rPr.set("spc", str(int(para.spacing * 100)))
    if para.caps:
        rPr.set("cap", "all")


def fill_text(tf, box: TextBox, lang: str, *, inherit: bool = False, zero_insets: bool = True) -> None:
    _body_pr(tf, box.anchor, zero_insets)
    # remove all but the first paragraph
    txBody = tf._txBody
    for extra in txBody.findall(qn("a:p"))[1:]:
        txBody.remove(extra)
    first = tf.paragraphs[0]
    for r in list(first._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            first._p.remove(r)
    for i, para in enumerate(box.paras):
        p = first if i == 0 else tf.add_paragraph()
        if not inherit:
            p.alignment = _ALIGN.get(para.align, PP_ALIGN.LEFT)
        if para.line:
            p.line_spacing = Pt(para.line)
        p.space_before = Pt(para.space_before)
        p.space_after = Pt(0)
        if para.level and inherit:
            p.level = min(para.level, 8)
        pPr = p._p.get_or_add_pPr()
        _insert_bullet_xml(pPr, para, inherit)
        if box.name == "slidenum" and not inherit:
            fld = etree.SubElement(p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
            rpr = etree.SubElement(fld, qn("a:rPr"), lang=lang, sz=str(int(para.size * 100)), b="1" if para.bold else "0")
            fill = etree.SubElement(rpr, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), val=para.color.lstrip("#").upper())
            etree.SubElement(rpr, qn("a:latin"), typeface=para.family)
            etree.SubElement(fld, qn("a:t")).text = para.text
            continue
        if para.lead and para.lead < len(para.text):
            _run(p, para.text[: para.lead], para, True, lang, inherit)
            _run(p, para.text[para.lead:], para, para.bold, lang, inherit)
        else:
            _run(p, para.text, para, para.bold, lang, inherit)


# --------------------------------------------------------------------------- shapes


def _add_shape(slide, s: Shape) -> None:
    sp = slide.shapes.add_shape(_SHAPES.get(s.kind, MSO_SHAPE.RECTANGLE), Emu(s.x), Emu(s.y), Emu(s.w), Emu(s.h))
    _strip_style(sp)
    if s.name:
        sp.name = s.name
    if s.fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(s.fill)
    else:
        sp.fill.background()
    if s.line:
        sp.line.color.rgb = _rgb(s.line)
        sp.line.width = Pt(s.line_w)
    else:
        sp.line.fill.background()
    if s.kind == "round" and s.radius:
        sp.adjustments[0] = min(0.5, emu(s.radius) / max(1, min(s.w, s.h)))


def _add_line(slide, ln: Line) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(ln.x1), Emu(ln.y1), Emu(ln.x2), Emu(ln.y2))
    _strip_style(c)
    c.line.color.rgb = _rgb(ln.color)
    c.line.width = Pt(ln.width)
    if ln.dash:
        from pptx.enum.dml import MSO_LINE

        c.line.dash_style = MSO_LINE.DASH


def _add_textbox(slide, box: TextBox, lang: str) -> None:
    tb = slide.shapes.add_textbox(Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    if box.name:
        tb.name = box.name
    fill_text(tb.text_frame, box, lang)


# --------------------------------------------------------------------------- table


def _cell_border(cell, bottom: str | None, width: float = 0.75) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    els = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.Element(qn(tag), w=str(emu(width)))
        if tag == "a:lnB" and bottom:
            f = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(f, qn("a:srgbClr"), val=bottom.lstrip("#").upper())
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        els.append(ln)
    for el in reversed(els):
        tcPr.insert(0, el)


def _add_table(slide, t: TableP, lang: str) -> None:
    nrows, ncols = len(t.rows) + 1, len(t.header)
    gf = slide.shapes.add_table(nrows, ncols, Emu(t.x), Emu(t.y), Emu(t.w), Emu(t.h))
    gf.name = "Tabel"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    for i, w in enumerate(t.col_w):
        tbl.columns[i].width = Emu(w)
    for j, h in enumerate(t.row_h):
        tbl.rows[j].height = Emu(h)
    pad = Emu(emu(t.pad))
    for j in range(nrows):
        values = t.header if j == 0 else t.rows[j - 1]
        for c in range(ncols):
            cell = tbl.cell(j, c)
            cell.margin_left = cell.margin_right = pad
            cell.margin_top = cell.margin_bottom = pad
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if j == 0:
                fill = t.header_fill
            elif j % 2 == 0:
                fill = t.zebra
            else:
                fill = t.background
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            _cell_border(cell, None if j == 0 or j == nrows - 1 else t.line)
            if j == nrows - 1:
                _cell_border(cell, t.header_fill, 1.5)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (t.numeric and t.numeric[c] and j > 0) else PP_ALIGN.LEFT
            p.line_spacing = Pt(round(t.size * 1.2, 2))
            r = p.add_run()
            r.text = values[c] if c < len(values) else ""
            r.font.size = Pt(t.size)
            r.font.name = t.header_family if j == 0 else t.family
            r.font.bold = j == 0
            r.font.color.rgb = _rgb(t.header_text if j == 0 else t.text)
            r._r.get_or_add_rPr().set("lang", lang)


# --------------------------------------------------------------------------- chart


def _number_format(values: list[float | None], unit: str) -> str:
    nums = [v for v in values if v is not None]
    decimals = 0 if all(float(v).is_integer() for v in nums) else 1
    base = "#,##0" + ("." + "0" * decimals if decimals else "")
    if unit.strip() == "%":
        return base + '"%"'
    return base


def _no_fill_xml(parent_el) -> None:
    """Transparent background + no border for chartSpace / plotArea."""
    spPr = parent_el.find(qn("c:spPr"))
    if spPr is None:
        spPr = etree.SubElement(parent_el, qn("c:spPr"))
        # c:spPr must precede c:txPr / externalData etc. in chartSpace; move into place
        for tag in ("c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst"):
            succ = parent_el.find(qn(tag))
            if succ is not None:
                succ.addprevious(spPr)
                break
    for child in list(spPr):
        spPr.remove(child)
    etree.SubElement(spPr, qn("a:noFill"))
    ln = etree.SubElement(spPr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def _add_chart(slide, cp: ChartP) -> None:
    c = cp.chart
    data = CategoryChartData()
    data.categories = c.categories
    all_vals = [v for s in c.series for v in s.values]
    fmt = _number_format(all_vals, c.unit)
    for s in c.series:
        data.add_series(s.name, s.values, number_format=fmt)
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.DOUGHNUT,
            "stacked": XL_CHART_TYPE.COLUMN_STACKED}[c.kind]
    gf = slide.shapes.add_chart(kind, Emu(cp.x), Emu(cp.y), Emu(cp.w), Emu(cp.h), data)
    gf.name = "Grafiek"
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(cp.size)
    ch.font.name = cp.family
    ch.font.color.rgb = _rgb(cp.muted)
    _no_fill_xml(ch._chartSpace)
    _no_fill_xml(ch._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")))
    rc = ch._chartSpace.find(qn("c:roundedCorners"))
    if rc is None:
        rc = etree.Element(qn("c:roundedCorners"))
        prev = None
        for tag in ("c:lang", "c:date1904"):
            prev = ch._chartSpace.find(qn(tag))
            if prev is not None:
                break
        if prev is not None:
            prev.addnext(rc)
        else:
            ch._chartSpace.insert(0, rc)
    rc.set("val", "0")
    multi = len(c.series) > 1
    plot = ch.plots[0]
    if c.kind == "pie":
        plot.vary_by_categories = True
        ser = plot.series[0]
        for i, _ in enumerate(c.categories):
            pt_ = ser.points[i]
            pt_.format.fill.solid()
            pt_.format.fill.fore_color.rgb = _rgb(cp.palette[i % len(cp.palette)])
            pt_.format.line.color.rgb = _rgb(cp.background)
            pt_.format.line.width = Pt(1.5)
        hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
        if hole is not None:
            hole.set("val", "58")
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage = True
        dl.show_value = False
        dl.show_category_name = False
        dl.number_format = "0%"
        dl.number_format_is_linked = False
        dl.font.size = Pt(cp.size)
        dl.font.bold = True
        dl.font.color.rgb = _rgb("#FFFFFF")
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(cp.size)
        ch.legend.font.color.rgb = _rgb(cp.text)
        return
    if c.kind in ("column", "bar", "stacked"):
        plot.gap_width = 60 if multi else 80
        plot.overlap = 100 if c.kind == "stacked" else (-8 if multi else 0)
    for i, ser in enumerate(plot.series):
        col = _rgb(cp.palette[i % len(cp.palette)])
        if c.kind == "line":
            ser.format.line.color.rgb = col
            ser.format.line.width = Pt(2.75)
            ser.smooth = False
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 7
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = col
            ser.marker.format.line.color.rgb = _rgb(cp.background)
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = col
            ser.format.line.fill.background()
    ca, va = ch.category_axis, ch.value_axis
    ca.format.line.color.rgb = _rgb(cp.grid)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(cp.size)
    ca.tick_labels.font.color.rgb = _rgb(cp.text)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(cp.size - 1)
    va.tick_labels.font.color.rgb = _rgb(cp.muted)
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    labels = (not multi or c.kind == "stacked") and len(c.categories) <= 12 and c.kind != "line"
    if labels and c.kind != "stacked":
        # Clean "label on the bar" style: no value axis, no gridlines.
        va.visible = False
        va.has_major_gridlines = False
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(cp.size)
        dl.font.bold = True
        dl.font.color.rgb = _rgb(cp.text)
    else:
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb(cp.grid)
        va.major_gridlines.format.line.width = Pt(0.75)
    if multi:
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(cp.size)
        ch.legend.font.color.rgb = _rgb(cp.text)
    else:
        ch.has_legend = False


# --------------------------------------------------------------------------- slides


def _picture(slide, p: Picture) -> None:
    pic = slide.shapes.add_picture(io.BytesIO(p.blob), Emu(p.x), Emu(p.y), Emu(p.w), Emu(p.h))
    if p.name:
        pic.name = p.name[:60]
        pic._element.nvPicPr.cNvPr.set("descr", p.name)


def _notes(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def remove_slides(prs) -> None:
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def write_deck(scenes: list[SlideScene], deck: DeckSpec, theme: Theme, out, *,
               tpl: "TemplateInfo | None" = None) -> None:
    lang = "nl-NL" if deck.language.startswith("nl") else ("en-GB" if deck.language.startswith("en") else deck.language)
    if tpl is not None:
        prs = Presentation(tpl.path)
        remove_slides(prs)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
        apply_theme_part(prs, theme)
    for sc in scenes:
        if tpl is not None:
            layout = tpl.layout_object(prs, sc.layout)
        else:
            layout = prs.slide_layouts[5]  # "Title Only": real title placeholder for outline/accessibility
        slide = prs.slides.add_slide(layout)
        if tpl is None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(sc.background)
        used: list = []  # keep element references (lxml proxies are only stable while referenced)
        for prim in sc.prims:
            if isinstance(prim, TextBox) and prim.placeholder:
                ph = _placeholder(slide, prim.placeholder, tpl)
                if ph is not None:
                    used.append(ph._element)
                    inherit = tpl is not None
                    x, y, w, h = prim.x, prim.y, prim.w, prim.h
                    if inherit:  # keep the template's insets; our rect is the inner text area
                        x, y, w, h = x - _PH_INSET_X, y - _PH_INSET_Y, w + 2 * _PH_INSET_X, h + 2 * _PH_INSET_Y
                    ph.left, ph.top, ph.width, ph.height = Emu(x), Emu(y), Emu(w), Emu(h)
                    fill_text(ph.text_frame, prim, lang, inherit=inherit, zero_insets=not inherit)
                    continue
            if isinstance(prim, Shape):
                _add_shape(slide, prim)
            elif isinstance(prim, Line):
                _add_line(slide, prim)
            elif isinstance(prim, TextBox):
                _add_textbox(slide, prim, lang)
            elif isinstance(prim, Picture):
                _picture(slide, prim)
            elif isinstance(prim, TableP):
                _add_table(slide, prim, lang)
            elif isinstance(prim, ChartP):
                _add_chart(slide, prim)
        # never leave empty placeholders behind
        for ph in list(slide.placeholders):
            if not any(ph._element is u for u in used):
                ph._element.getparent().remove(ph._element)
        if tpl is not None and sc.meta.get("slidenum"):
            tpl.add_slide_number(slide, layout)
        _notes(slide, sc.notes)
    _register_notes_master(prs)
    cp = prs.core_properties
    cp.title = deck.title[:250]
    cp.author = deck.author[:250] or "DocWizard"
    cp.subject = deck.subtitle[:250]
    cp.language = lang
    cp.keywords = deck.audience[:250]
    cp.last_modified_by = "DocWizard"
    cp.revision = 1
    prs.save(out)


def _register_notes_master(prs) -> None:
    """python-pptx relates the notes master but omits <p:notesMasterIdLst>; strict readers (macOS Quick Look,
    some viewers) then fail to open the file. Add the list when missing."""
    pres = prs.part._element
    if pres.find(qn("p:notesMasterIdLst")) is not None:
        return
    rid = next((r.rId for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER), None)
    if rid is None:
        return
    lst = etree.Element(qn("p:notesMasterIdLst"))
    etree.SubElement(lst, qn("p:notesMasterId")).set(qn("r:id"), rid)
    anchor = pres.find(qn("p:sldMasterIdLst"))
    anchor.addnext(lst)


def _placeholder(slide, role: str, tpl: "TemplateInfo | None"):
    from pptx.enum.shapes import PP_PLACEHOLDER as PH

    phs = list(slide.placeholders)
    if role == "title":
        for ph in phs:
            if ph.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE):
                return ph
        return None
    if role == "subtitle":
        for ph in phs:
            if ph.placeholder_format.type == PH.SUBTITLE:
                return ph
        role = "body"
    bodies = [ph for ph in phs if ph.placeholder_format.type in (PH.BODY, PH.OBJECT)]
    bodies.sort(key=lambda p: (p.left or 0, p.top or 0))
    if role == "body":
        return bodies[0] if bodies else None
    if role == "body2":
        return bodies[1] if len(bodies) > 1 else None
    return None


__all__ = ["write_deck", "apply_theme_part", "fill_text", "remove_slides"]
