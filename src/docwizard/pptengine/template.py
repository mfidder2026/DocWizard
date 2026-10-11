"""User PowerPoint templates: analysis, layout selection and placeholder geometry.

`TemplateInfo` reads a .pptx/.potx once (slide size, theme colours/fonts, master text styles, layouts and
their placeholders, decorative master/layout art for previews) and answers the questions the layout engine
asks: which layout serves a role, where its title/body placeholders are, what size/colour/alignment the
template's title style has, and where free content may be drawn.
"""

from __future__ import annotations

import copy
import re
import zipfile
from functools import lru_cache
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .layout import Area
from .scene import Picture, Prim, Shape, emu
from .themes import Theme, _scheme_colours, _scheme_fonts, theme_from_colours

ROLES = ("title", "section", "content", "two", "title_only", "blank")
_FALLBACK = {
    "title": ("title", "section", "title_only", "blank", "content"),
    "section": ("section", "title_only", "blank", "content"),
    "content": ("content", "title_only", "blank"),
    "two": ("two", "title_only", "blank", "content"),
    "title_only": ("title_only", "content", "blank", "section"),
    "blank": ("blank", "title_only", "content"),
}
_TYPE_ROLE = {"title": "title", "secHead": "section", "obj": "content", "tx": "content", "twoObj": "two",
              "twoTxTwoObj": "two_alt", "titleOnly": "title_only", "blank": "blank", "objTx": "content_alt"}
_SCHEME_MAP = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_INSET_X, _INSET_Y = 7.2, 3.6
_BODY_TYPES = (PH.BODY, PH.OBJECT)
_TITLE_TYPES = (PH.TITLE, PH.CENTER_TITLE)
_FOOTER_TYPES = (PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER)


def _pt(v: int | None) -> float:
    return (v or 0) / 12700


def _area(shape) -> Area | None:
    try:
        if shape.left is None or shape.width is None:
            return None
        return Area(_pt(shape.left), _pt(shape.top), _pt(shape.width), _pt(shape.height))
    except Exception:
        return None


def _inner(a: Area | None) -> Area | None:
    if a is None:
        return None
    return Area(a.x + _INSET_X, a.y + _INSET_Y, max(10.0, a.w - 2 * _INSET_X), max(10.0, a.h - 2 * _INSET_Y))


class TemplateInfo:
    def __init__(self, path: str):
        self.path = str(path)
        prs = Presentation(self.path)
        self.W, self.H = _pt(prs.slide_width), _pt(prs.slide_height)
        master = prs.slide_masters[0]
        theme_xml = master.part.part_related_by(RT.THEME).blob
        self.colours = _scheme_colours(theme_xml)
        self.major, self.minor = _scheme_fonts(theme_xml)
        self.existing_slides = len(prs.slides)
        self.masters = len(prs.slide_masters)
        self._master_bg = self._bg_colour(master._element)
        bg = self._master_bg or self.colours.get("lt1", "#FFFFFF")
        accent = self.colours.get("accent1") or self.colours.get("dk2") or "#1F4E9A"
        self.theme: Theme = theme_from_colours(
            "sjabloon", "PowerPoint-sjabloon", accent=accent, colours=self.colours,
            heading_font=self.major or "Calibri", body_font=self.minor or "Calibri", background=bg)
        # master text styles
        tx = master._element.find(qn("p:txStyles"))
        self._master_title = self._lvl_style(tx.find(qn("p:titleStyle")) if tx is not None else None)
        self._master_body = self._lvl_style(tx.find(qn("p:bodyStyle")) if tx is not None else None)
        self._master_title_anchor = "m"
        self.master_body: Area | None = None
        self.footer_top: float | None = None
        for ph in master.placeholders:
            t = ph.placeholder_format.type
            if t in _TITLE_TYPES:
                self._master_title_anchor = self._anchor(ph._element) or "m"
            elif t in _BODY_TYPES and self.master_body is None:
                self.master_body = _area(ph)
            elif t in _FOOTER_TYPES:
                a = _area(ph)
                if a and a.y > self.H * 0.6:
                    self.footer_top = a.y if self.footer_top is None else min(self.footer_top, a.y)
        # layouts
        self.layouts: list[dict[str, Any]] = []
        flat = 0
        for mi, m in enumerate(prs.slide_masters):
            for li, lay in enumerate(m.slide_layouts):
                phs = []
                for ph in lay.placeholders:
                    pf = ph.placeholder_format
                    phs.append({"type": pf.type, "idx": pf.idx, "name": ph.name, "area": _area(ph),
                                "el": ph._element})
                self.layouts.append({"index": flat, "master": mi, "layout": li, "name": lay.name,
                                     "type": lay._element.get("type", ""), "phs": phs,
                                     "show_master": lay._element.get("showMasterSp", "1") not in ("0", "false"),
                                     "bg": self._bg_colour(lay._element)})
                flat += 1
        self.roles = self._assign_roles()
        self._decor: dict[str, list[Prim]] = {}
        self._prs = prs
        # styles used by the layout engine
        ts = self.style("title_only", "title")
        self.title_size, self.title_color, self.title_align, self.title_anchor = (
            ts["size"], ts["color"], ts["align"], ts["anchor"])
        self.title_font = ts["font"]
        cs = self.style("title", "title")
        self.ctr_title_size, self.ctr_title_align, self.ctr_title_anchor = cs["size"], cs["align"], cs["anchor"]
        ss = self.style("title", "subtitle")
        self.subtitle_size, self.subtitle_anchor = ss["size"], ss["anchor"]
        self.subtitle_align = ss["align"]
        sec = self.style("section", "title")
        self.section_align, self.section_anchor = sec["align"], sec["anchor"]
        self.body_size = self.style("content", "body")["size"]

    # ----------------------------------------------------------------- parsing helpers

    def _colour_of(self, el) -> str:
        if el is None:
            return ""
        s = el.find(qn("a:srgbClr"))
        if s is not None:
            return "#" + s.get("val", "000000").upper()
        s = el.find(qn("a:schemeClr"))
        if s is not None:
            v = s.get("val", "")
            return self.colours.get(_SCHEME_MAP.get(v, v), "")
        return ""

    def _bg_colour(self, root) -> str:
        bg = root.find(f"{qn('p:cSld')}/{qn('p:bg')}")
        if bg is None:
            return ""
        pr = bg.find(qn("p:bgPr"))
        if pr is not None:
            return self._colour_of(pr.find(qn("a:solidFill")))
        ref = bg.find(qn("p:bgRef"))
        if ref is not None:
            return self._colour_of(ref)
        return ""

    def _lvl_style(self, el) -> dict[str, Any]:
        out: dict[str, Any] = {}
        if el is None:
            return out
        lvl = el.find(qn("a:lvl1pPr"))
        if lvl is None:
            return out
        if lvl.get("algn"):
            out["align"] = {"ctr": "c", "r": "r"}.get(lvl.get("algn"), "l")
        d = lvl.find(qn("a:defRPr"))
        if d is not None:
            if d.get("sz"):
                out["size"] = int(d.get("sz")) / 100
            col = self._colour_of(d.find(qn("a:solidFill")))
            if col:
                out["color"] = col
            lat = d.find(qn("a:latin"))
            if lat is not None and lat.get("typeface"):
                tf = lat.get("typeface")
                out["font"] = self.major if tf == "+mj-lt" else self.minor if tf == "+mn-lt" else tf
        return out

    @staticmethod
    def _anchor(sp_el) -> str:
        bp = sp_el.find(f".//{qn('a:bodyPr')}")
        if bp is None or not bp.get("anchor"):
            return ""
        return {"t": "t", "ctr": "m", "b": "b"}.get(bp.get("anchor"), "t")

    # ----------------------------------------------------------------- roles

    def _infer_role(self, lay: dict[str, Any]) -> str:
        types = [p["type"] for p in lay["phs"] if p["type"] not in _FOOTER_TYPES]
        name = lay["name"].lower()
        has_title = any(t in _TITLE_TYPES for t in types)
        bodies = [t for t in types if t in _BODY_TYPES]
        if PH.CENTER_TITLE in types or re.search(r"\btitel(dia)?\b|title slide|cover|voorblad", name):
            return "title"
        if re.search(r"sect|hoofdstuk|tussen|divider|chapter", name):
            return "section"
        if has_title and len(bodies) >= 2:
            return "two"
        if has_title and len(bodies) == 1:
            return "content"
        if has_title and not bodies and len(types) == 1:
            return "title_only"
        if not types:
            return "blank"
        return "other"

    def _assign_roles(self) -> dict[str, int]:
        roles: dict[str, int] = {}
        for lay in self.layouts:  # explicit layout types first
            r = _TYPE_ROLE.get(lay["type"], "")
            if r in ROLES and r not in roles and self._usable(lay, r):
                roles[r] = lay["index"]
        for lay in self.layouts:
            r = _TYPE_ROLE.get(lay["type"], "")
            if r == "two_alt" and "two" not in roles and self._usable(lay, "two"):
                roles["two"] = lay["index"]
            if r == "content_alt" and "content" not in roles and self._usable(lay, "content"):
                roles["content"] = lay["index"]
        for lay in self.layouts:  # custom layouts
            r = self._infer_role(lay)
            if r in ROLES and r not in roles and self._usable(lay, r):
                roles[r] = lay["index"]
        return roles

    def _usable(self, lay: dict[str, Any], role: str) -> bool:
        types = [p["type"] for p in lay["phs"]]
        has_title = any(t in _TITLE_TYPES for t in types)
        bodies = [p for p in lay["phs"] if p["type"] in _BODY_TYPES and p["area"] is not None]
        if role in ("title", "section", "title_only", "content", "two") and not has_title:
            return False
        if role == "content":
            return len(bodies) >= 1
        if role == "two":
            return len(bodies) >= 2
        return True

    def resolve_role(self, role: str) -> str:
        for r in _FALLBACK.get(role, (role,)):
            if r in self.roles:
                return r
        return next(iter(self.roles), "blank")

    def _layout(self, role: str) -> dict[str, Any] | None:
        r = self.resolve_role(role)
        idx = self.roles.get(r)
        return self.layouts[idx] if idx is not None else None

    def layout_object(self, prs, role: str):
        lay = self._layout(role)
        if lay is None:
            return prs.slide_layouts[-1]
        return prs.slide_masters[lay["master"]].slide_layouts[lay["layout"]]

    # ----------------------------------------------------------------- geometry

    def _ph(self, role: str, key: str) -> dict[str, Any] | None:
        lay = self._layout(role)
        if lay is None:
            return None
        phs = [p for p in lay["phs"] if p["area"] is not None]
        if key == "title":
            return next((p for p in phs if p["type"] in _TITLE_TYPES), None)
        if key == "subtitle":
            sub = next((p for p in phs if p["type"] == PH.SUBTITLE), None)
            if sub is not None:
                return sub
            key = "body"
        bodies = sorted((p for p in phs if p["type"] in _BODY_TYPES), key=lambda p: (p["area"].x, p["area"].y))
        if key == "body":
            return bodies[0] if bodies else None
        if key == "body2":
            return bodies[1] if len(bodies) > 1 else None
        return None

    def ph_rect(self, role: str, key: str) -> Area | None:
        p = self._ph(role, key)
        return _inner(p["area"]) if p else None

    def content_rect(self, role: str) -> Area:
        r = self.resolve_role(role)
        title = self.ph_rect(r, "title")
        if r == "content":
            b = self.ph_rect(r, "body")
            if b is not None:
                return b
        if r == "two":
            b1, b2 = self.ph_rect(r, "body"), self.ph_rect(r, "body2")
            if b1 and b2:
                return Area(b1.x, min(b1.y, b2.y), b2.r - b1.x, max(b1.b, b2.b) - min(b1.y, b2.y))
        base = _inner(self.master_body) if self.master_body else None
        if base is None:
            mx = self.W * 0.06
            base = Area(mx, self.H * 0.22, self.W - 2 * mx, self.H * 0.68)
        top = base.y
        if title is not None:
            top = max(top, title.b + 14)
        bottom = base.b
        if self.footer_top is not None and self.footer_top > top + 60:
            bottom = min(bottom, self.footer_top - 6)
        if title is not None and r in ("title_only", "blank"):
            x = min(base.x, title.x) if abs(base.x - title.x) < 30 else base.x
            return Area(x, top, base.r - x, max(60.0, bottom - top))
        return Area(base.x, top, base.w, max(60.0, bottom - top))

    def style(self, role: str, key: str) -> dict[str, Any]:
        """Effective size/colour/alignment/anchor/font of a layout placeholder (layout over master)."""
        is_title = key == "title"
        m = self._master_title if is_title else self._master_body
        out = {"size": m.get("size", 44.0 if is_title else 28.0), "color": m.get("color", ""),
               "align": m.get("align", "l"), "anchor": self._master_title_anchor if is_title else "t",
               "font": m.get("font", self.major if is_title else self.minor)}
        p = self._ph(role, key)
        if p is not None:
            el = p["el"]
            lst = el.find(f".//{qn('a:lstStyle')}")
            if lst is not None:
                out.update({k: v for k, v in self._lvl_style(lst).items() if v})
            a = self._anchor(el)
            if a:
                out["anchor"] = a
        if not out["color"]:
            out["color"] = self.colours.get("dk1", "#000000") if not self.theme.dark else self.theme.text
        return out

    def background_for(self, role: str) -> str:
        lay = self._layout(role)
        return (lay or {}).get("bg") or self._master_bg or self.colours.get("lt1", "#FFFFFF")

    @property
    def background(self) -> str:
        return self._master_bg or self.colours.get("lt1", "#FFFFFF")

    # ----------------------------------------------------------------- decor (previews)

    def decor(self, role: str) -> list[Prim]:
        r = self.resolve_role(role)
        if r in self._decor:
            return self._decor[r]
        out: list[Prim] = []
        lay = self._layout(r)
        if lay is not None:
            lobj = self._prs.slide_masters[lay["master"]].slide_layouts[lay["layout"]]
            bg = self.background_for(r)
            out.append(Shape("rect", 0, 0, emu(self.W), emu(self.H), bg))
            shapes = (list(lobj.slide_master.shapes) if lay["show_master"] else []) + list(lobj.shapes)
            for sh in shapes:
                out.extend(self._shape_prims(sh))
        self._decor[r] = out
        return out

    def _shape_prims(self, sh) -> list[Prim]:
        try:
            if sh.is_placeholder:
                return []
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                return [p for s in sh.shapes for p in self._shape_prims(s)]
            if sh.left is None:
                return []
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return [Picture(sh.left, sh.top, sh.width, sh.height, sh.image.blob, sh.image.ext)]
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                sp_pr = sh._element.find(qn("p:spPr"))
                fill = self._colour_of(sp_pr.find(qn("a:solidFill"))) if sp_pr is not None else ""
                if fill:
                    return [Shape("rect", sh.left, sh.top, sh.width, sh.height, fill)]
        except Exception:
            return []
        return []

    # ----------------------------------------------------------------- writing helpers

    def add_slide_number(self, slide, layout) -> None:
        for ph in layout.placeholders:
            if ph.placeholder_format.type == PH.SLIDE_NUMBER:
                slide.shapes._spTree.append(copy.deepcopy(ph._element))
                return

    def profile(self) -> dict[str, Any]:
        def ph_info(p: dict[str, Any]) -> dict[str, Any]:
            a = p["area"]
            return {"type": str(p["type"]).split(" ")[0].replace("PP_PLACEHOLDER.", "").lower()
                    if p["type"] is not None else "", "idx": p["idx"], "name": p["name"],
                    "left": emu(a.x) if a else None, "top": emu(a.y) if a else None,
                    "width": emu(a.w) if a else None, "height": emu(a.h) if a else None}

        role_of = {v: k for k, v in self.roles.items()}
        return {
            "path": self.path,
            "slide_width": emu(self.W), "slide_height": emu(self.H),
            "width_in": round(self.W / 72, 3), "height_in": round(self.H / 72, 3),
            "aspect": "16:9" if abs(self.W / self.H - 16 / 9) < 0.02 else (
                "4:3" if abs(self.W / self.H - 4 / 3) < 0.02 else f"{self.W / self.H:.2f}"),
            "theme": {"colors": dict(self.colours), "fonts": {"major": self.major, "minor": self.minor}},
            "layouts": [{"index": lay["index"], "name": lay["name"], "type": lay["type"],
                         "role": role_of.get(lay["index"], ""),
                         "placeholders": [ph_info(p) for p in lay["phs"]]} for lay in self.layouts],
            "roles": {r: self.layouts[i]["name"] for r, i in self.roles.items()},
            "masters": self.masters,
            "slides": self.existing_slides,
            "title_style": {"size": self.title_size, "color": self.title_color, "align": self.title_align,
                            "font": self.title_font},
            "body_size": self.body_size,
        }


@lru_cache(maxsize=8)
def _cached(path: str, mtime: float) -> TemplateInfo:
    return TemplateInfo(path)


def load_template(path: str) -> TemplateInfo:
    import os

    p = str(path)
    if p.lower().endswith(".potx"):
        return TemplateInfo(_potx_to_pptx(p))
    return _cached(p, os.path.getmtime(p))


def _potx_to_pptx(path: str) -> str:
    """python-pptx refuses .potx content types; rewrite the main part's content type in a temp copy."""
    import tempfile

    out = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    out.close()
    return out.name


def analyze_pptx_template(path: str) -> dict[str, Any]:
    """Slide size, theme colours/fonts, layouts with placeholders (type, idx, position) and role mapping."""
    return load_template(path).profile()

