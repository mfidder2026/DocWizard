"""Tests for the deterministic PowerPoint engine (docwizard.pptengine)."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from docwizard.pptengine import (
    DECK_JSON_SCHEMA,
    DECK_PROMPT_GUIDE,
    SLIDE_TYPES,
    THEMES,
    DeckSpec,
    analyze_pptx_template,
    build_pptx,
    normalize_deck,
    plan,
    render_previews,
    theme_from_docx,
)
from docwizard.pptengine.render import _para_lines
from docwizard.pptengine.scene import EMU_PER_PT, TextBox
from docwizard.pptengine.spec import fmt_number, parse_number

FIXTURE_DOCX = Path(__file__).parent / "fixtures" / "aqs_template.docx"


def _png(color=(30, 80, 160), size=(800, 450)) -> bytes:
    im = Image.new("RGB", size, (255, 255, 255))
    im.paste(color, (100, 100, 700, 350))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def images(ref: str):
    return (_png(), "png") if ref.startswith("diagram:") or ref == "foto.png" else None


FULL = {
    "title": "Centrale wachtwoordkluis voor beheeraccounts",
    "subtitle": "Advies aan de stuurgroep",
    "author": "Enterprise architectuur",
    "date": "oktober 2026",
    "language": "nl",
    "audience": "Stuurgroep IV",
    "slides": [
        {"type": "title", "title": "Centrale wachtwoordkluis voor beheeraccounts"},
        {"type": "bullets", "title": "Wij adviseren een centrale kluis vóór Q3 2027",
         "bullets": ["Risico: accounts zijn verspreid", {"text": "Kosten: terugverdiend in twee jaar",
                                                         "sub": ["Minder incidenten"]}],
         "notes": "Eerst het advies."},
        {"type": "section", "title": "Waarom nu"},
        {"type": "kpi", "title": "Het beheer kost tijd en vergroot het risico",
         "kpis": [{"value": "1.240", "label": "Beheeraccounts"}, {"value": "62%", "label": "Gedeeld"},
                  {"value": "€ 1,2 mln", "label": "Schade per incident", "detail": "sectorcijfers"}],
         "source": "Inventarisatie 2026"},
        {"type": "quote", "title": "De auditor is kritisch",
         "quote": {"text": "Wij kunnen de beheersing niet bevestigen.", "source": "Externe auditor"}},
        {"type": "chart", "title": "Incidenten stijgen al drie jaar",
         "chart": {"kind": "column", "categories": ["2023", "2024", "2025"],
                   "series": [{"name": "Incidenten", "values": [19, 27, 33]}]}, "bullets": ["Verdubbeling"]},
        {"type": "chart", "title": "Kosten dalen, besparing stijgt",
         "chart": {"kind": "line", "categories": ["2027", "2028"],
                   "series": [{"name": "Kosten", "values": [0.4, 0.3]}, {"name": "Besparing", "values": [0.1, 0.5]}],
                   "unit": "€ mln"}},
        {"type": "chart", "title": "Verdeling", "chart": {"kind": "pie", "categories": ["A", "B", "C"],
                                                          "series": [{"name": "x", "values": [50, 30, 20]}]}},
        {"type": "chart", "title": "Gestapeld", "chart": {"kind": "stacked", "categories": ["A", "B"],
                                                          "series": [{"name": "x", "values": [5, 3]},
                                                                     {"name": "y", "values": [2, 4]}]}},
        {"type": "chart", "title": "Liggend", "chart": {"kind": "bar", "categories": ["A", "B"],
                                                        "series": [{"name": "x", "values": [5, 3]}]}},
        {"type": "section", "title": "Wat we voorstellen"},
        {"type": "diagram", "title": "De kluis wordt de enige toegangsweg",
         "image": {"ref": "diagram:kluis", "caption": "Doelarchitectuur"}},
        {"type": "image", "title": "Zo ziet het scherm eruit", "image": {"ref": "foto.png"}},
        {"type": "comparison", "title": "De kluis scoort beter",
         "comparison": {"left": {"heading": "Nu", "bullets": ["Excel"]},
                        "right": {"heading": "Kluis", "bullets": ["Rotatie"]}, "verdict": "Kies de kluis"}},
        {"type": "two_column", "title": "Techniek en organisatie",
         "columns": [{"heading": "Techniek", "bullets": ["Hoge beschikbaarheid"]},
                     {"heading": "Organisatie", "bullets": ["Training"]}]},
        {"type": "table", "title": "Optie B biedt de beste prijs-kwaliteit",
         "table": {"header": ["Criterium", "A", "B"], "rows": [["Kosten", "€ 180.000", "€ 140.000"],
                                                                ["Referenties", "4", "7"]]}},
        {"type": "timeline", "title": "In vier fasen klaar",
         "steps": [{"label": "Aanbesteding", "date": "Q4 2026"}, {"label": "Pilot", "date": "Q1 2027"},
                   {"label": "Uitrol", "date": "Q2 2027", "text": "Alle systemen"}]},
        {"type": "closing", "title": "Wij vragen drie besluiten", "bullets": ["Akkoord", "Budget", "Start"]},
    ],
}


# --------------------------------------------------------------------------- normalisation


def test_parse_number_dutch_formats():
    assert parse_number("€ 1,2 mln") == (1.2, "€ mln")
    assert parse_number("45%") == (45.0, "%")
    assert parse_number("1.234") == (1234.0, "")
    assert parse_number("1.234,5") == (1234.5, "")
    assert parse_number("3.5") == (3.5, "")
    assert parse_number(7) == (7.0, "")
    assert parse_number("geen") == (None, "")
    assert fmt_number(1234.5) == "1.234,5"


def test_normalize_messy_llm_output():
    raw = """Hier is de presentatie:
```json
{"titel": "Kluis", "taal": "nl", "dia's": [
  {"soort": "titel", "kop": "Kluis voor beheer"},
  {"layout": "opsomming", "titel": "Een **heel** belangrijk punt", "punten": "Eerste punt\\nTweede punt\\n  - sub van twee"},
  {"titel": "Cijfers", "cijfers": [{"waarde": 12, "label": "Systemen"}, "62% - gedeeld"]},
  {"type": "grafiek", "titel": "Groei", "grafiek": {"soort": "staaf", "labels": ["2024", "2025"],
     "waarden": ["€ 1,2 mln", "€ 2,4 mln"]}},
  {"type": "hologram", "title": "Raar type", "bullets": ["a", "b"]},
  {"type": "tijdlijn", "title": "Planning", "stappen": ["Q1 2027: pilot", "Q2 2027: uitrol"]},
  {"type": "vergelijking", "title": "A of B", "links": ["Optie A", "duur"], "rechts": ["Optie B", "goedkoop"],
   "advies": "Kies B"},
  {"type": "sectie", "title": "Deel twee"},
  {"type": "sectie", "title": "Deel drie"},
  {"type": "bullets", "title": "Te veel", "bullets": ["x"]}
]}
```"""
    deck = normalize_deck(raw)
    types = [s.type for s in deck.slides]
    assert types[0] == "title" and types[-1] == "closing"
    assert types[1] == "agenda"  # auto agenda: two sections
    assert [b.text for b in deck.slides[1].bullets] == ["Deel twee", "Deel drie"]
    assert deck.title == "Kluis"
    b = deck.slides[2]
    assert b.title == "Een heel belangrijk punt"
    assert [x.text for x in b.bullets] == ["Eerste punt", "Tweede punt"] and b.bullets[1].sub == ["sub van twee"]
    kpi = deck.slides[3]
    assert kpi.type == "kpi" and kpi.kpis[0].value == "12" and kpi.kpis[1].value == "62%"
    ch = deck.slides[4].chart
    assert ch is not None and ch.kind == "column" and ch.series[0].values == [1.2, 2.4] and ch.unit == "€ mln"
    assert deck.slides[5].type == "bullets"
    assert any("hologram" in w for w in deck.warnings)
    tl = deck.slides[6]
    assert tl.type == "timeline" and tl.steps[0].date == "Q1 2027" and tl.steps[0].label == "pilot"
    comp = deck.slides[7].comparison
    assert comp is not None and comp.left.heading == "Optie A" and comp.verdict == "Kies B"
    assert deck.slides[8].number == 1 and deck.slides[9].number == 2
    assert deck.slides[10].kicker == "Deel drie"


def test_normalize_list_input_limits_and_truncation():
    long_title = "woord " * 60
    deck = normalize_deck([{"title": long_title, "bullets": [f"punt {i}" for i in range(30)]}])
    s = deck.slides[1]
    assert len(s.title) <= 150 and s.title.endswith("…")
    assert len(s.bullets) == 14
    assert any("ingekort" in w for w in deck.warnings)
    assert deck.slides[-1].title == "Vragen?"


def test_normalize_trailing_comma_and_truncated_json():
    deck = normalize_deck('{"title": "X", "slides": [{"type": "bullets", "title": "A", "bullets": ["a",],},')
    assert deck.slides[1].title == "A"
    with pytest.raises(ValueError):
        normalize_deck("geen json")


# --------------------------------------------------------------------------- building


def _textbox_overflows(scenes) -> list[str]:
    bad = []
    for n, sc in enumerate(scenes, 1):
        for p in sc.prims:
            if not isinstance(p, TextBox):
                continue
            w = p.w / EMU_PER_PT
            h = sum(len(_para_lines(q, w)) * q.line + (q.space_before if i else 0) for i, q in enumerate(p.paras))
            if h > p.h / EMU_PER_PT + 0.5:
                bad.append(f"slide {n} {p.name or p.paras[0].text[:30]}: {h:.1f} > {p.h / EMU_PER_PT:.1f}")
            if p.x < 0 or p.y < 0 or p.x + p.w > 12192000 + 1 or p.y + p.h > 6858000 + 1:
                bad.append(f"slide {n}: textbox outside slide")
    return bad


@pytest.mark.parametrize("theme", list(THEMES))
def test_every_slide_type_builds_and_reopens(tmp_path, theme):
    out = tmp_path / f"{theme}.pptx"
    report = build_pptx(FULL, out, theme=theme, images=images)
    assert report["split"] == 0
    prs = Presentation(str(out))
    assert len(prs.slides) == report["slides"] == len(normalize_deck(FULL).slides)
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000
    assert prs.core_properties.title == FULL["title"]
    assert prs.core_properties.author == FULL["author"]
    kinds = {s.type for s in normalize_deck(FULL).slides}
    assert kinds == set(SLIDE_TYPES)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    tables = [sh for s in prs.slides for sh in s.shapes if sh.has_table]
    pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
    assert len(charts) == 5 and len(tables) == 1 and len(pics) == 2
    # every slide has a real title placeholder with text (outline view / accessibility)
    for s in prs.slides:
        assert s.shapes.title is not None and s.shapes.title.text_frame.text.strip()
        for ph in s.placeholders:
            assert ph.has_text_frame and ph.text_frame.text.strip()
    assert "Eerst het advies." in prs.slides[2].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(out) as z:
        assert b"notesMasterIdLst" in z.read("ppt/presentation.xml")
        theme_xml = z.read("ppt/theme/theme1.xml").decode()
        assert THEMES[theme].accent.lstrip("#").upper() in theme_xml or THEMES[theme].chart[0].lstrip("#") in theme_xml


@pytest.mark.parametrize("theme", list(THEMES))
def test_no_text_overflow(theme):
    _, scenes, report, _, _ = plan(FULL, theme=theme, images=images)
    assert _textbox_overflows(scenes) == []


def test_split_long_bullet_slide(tmp_path):
    bullets = [f"Dit is bullet nummer {i} met behoorlijk veel tekst om de dia echt te vullen en te laten overlopen"
               for i in range(14)]
    deck = {"title": "Split", "slides": [{"type": "bullets", "title": "Heel veel punten", "bullets": bullets}]}
    out = tmp_path / "split.pptx"
    report = build_pptx(deck, out)
    assert report["split"] >= 1
    assert any("gesplitst" in w for w in report["warnings"])
    prs = Presentation(str(out))
    titles = [s.shapes.title.text_frame.text for s in prs.slides]
    assert any(t.endswith("(vervolg)") for t in titles)
    body = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    for b in bullets:
        assert b in body  # nothing lost
    _, scenes, _, _, _ = plan(deck)
    assert _textbox_overflows(scenes) == []


def test_split_long_table_and_timeline():
    rows = [[f"Rij {i} met een langere omschrijving", f"{i * 1000}", "Ja"] for i in range(30)]
    steps = [{"label": f"Stap {i}", "date": f"Q{i % 4 + 1}"} for i in range(10)]
    deck = {"title": "T", "slides": [
        {"type": "table", "title": "Grote tabel", "table": {"header": ["Naam", "Bedrag", "OK"], "rows": rows}},
        {"type": "timeline", "title": "Lange tijdlijn", "steps": steps}]}
    _, scenes, report, _, _ = plan(deck)
    assert report["split"] >= 2
    assert _textbox_overflows(scenes) == []


def test_missing_image_uses_placeholder_with_warning(tmp_path):
    deck = {"title": "X", "slides": [{"type": "image", "title": "Foto", "image": {"ref": "bestaat-niet.png"}}]}
    report = build_pptx(deck, tmp_path / "x.pptx", images=images)
    assert any("niet gevonden" in w for w in report["warnings"])


def test_previews_are_pngs():
    pngs = render_previews(FULL, theme="donker", images=images, width=640)
    assert len(pngs) == len(normalize_deck(FULL).slides)
    for p in pngs:
        assert p[:8] == b"\x89PNG\r\n\x1a\n"
        im = Image.open(io.BytesIO(p))
        assert im.size == (640, 360)


# --------------------------------------------------------------------------- themes & templates


def test_theme_from_docx():
    th = theme_from_docx(str(FIXTURE_DOCX))
    assert th.accent == "#004B87"  # Heading 1 colour of the house style
    assert th.heading_font == "Calibri Light"
    assert th.body_font == "Calibri"
    assert len(th.chart) == 6
    pngs = render_previews(FULL, theme=th, images=images, width=320)
    assert len(pngs) > 10


def _make_template(path: Path) -> None:
    prs = Presentation()  # python-pptx default template (4:3), with one existing slide
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Bestaande dia"
    prs.save(str(path))


def test_analyze_template(tmp_path):
    tpl = tmp_path / "tpl.pptx"
    _make_template(tpl)
    prof = analyze_pptx_template(str(tpl))
    assert prof["aspect"] == "4:3" and prof["slides"] == 1
    assert {"title", "section", "content", "two", "title_only", "blank"} <= set(prof["roles"])
    assert prof["theme"]["fonts"]["major"] == "Calibri"
    lay = next(lay for lay in prof["layouts"] if lay["role"] == "content")
    assert any(p["type"] == "body" or p["type"] == "object" for p in lay["placeholders"])
    assert all(p["left"] is not None for p in lay["placeholders"])
    json.dumps(prof)


def test_template_mode(tmp_path):
    tpl = tmp_path / "tpl.pptx"
    _make_template(tpl)
    out = tmp_path / "out.pptx"
    report = build_pptx(FULL, out, template=str(tpl), images=images)
    prs = Presentation(str(out))
    assert len(prs.slides) == report["slides"]
    assert prs.slide_width == 9144000  # template size kept
    titles = [s.shapes.title.text_frame.text if s.shapes.title is not None else "" for s in prs.slides]
    assert "Bestaande dia" not in titles
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "Title Slide"
    assert "Title and Content" in layouts and "Section Header" in layouts and "Two Content" in layouts
    assert "Title Only" in layouts
    for s in prs.slides:
        for ph in s.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            assert ph.has_text_frame and ph.text_frame.text.strip(), (s.slide_layout.name, ph.name)
    bullets = next(s for s in prs.slides if s.slide_layout.name == "Title and Content")
    body = next(ph for ph in bullets.placeholders if ph.placeholder_format.idx == 1)
    assert "accounts zijn verspreid" in body.text_frame.text
    _, scenes, _, _, _ = plan(FULL, template=str(tpl), images=images)
    assert _textbox_overflows(scenes) == []
    pngs = render_previews(FULL, template=str(tpl), images=images, width=400)
    assert Image.open(io.BytesIO(pngs[0])).size == (400, 300)


# --------------------------------------------------------------------------- prompt


def test_prompt_guide_and_schema():
    assert "action title" in DECK_PROMPT_GUIDE.lower()
    assert "Pyramid" in DECK_PROMPT_GUIDE
    schema = json.loads(json.dumps(DECK_JSON_SCHEMA))
    assert schema["properties"]["slides"]["items"]["properties"]["type"]["enum"] == list(SLIDE_TYPES)
    example = DECK_PROMPT_GUIDE[DECK_PROMPT_GUIDE.index('{"title"'):]
    deck = normalize_deck(example)
    assert isinstance(deck, DeckSpec) and deck.slides[0].type == "title" and deck.slides[-1].type == "closing"
    assert not deck.warnings
