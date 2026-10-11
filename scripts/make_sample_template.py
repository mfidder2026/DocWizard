"""Build a realistic 'Architecture Quick Scan' house-style template for tests.

Mimics a Dutch Word install (localized style IDs like 'Kop1'), with a cover page
(content controls + placeholders), a TOC field, a section break, header/footer,
multi-level numbered headings, a custom body style, guidance text and a styled
table. Usage: uv run python scripts/make_sample_template.py out.docx
"""
import sys
from copy import deepcopy

import docx
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Cm

BRAND = RGBColor(0x00, 0x4B, 0x87)
out = sys.argv[1] if len(sys.argv) > 1 else "sample_aqs_template.docx"

d = docx.Document()
st = d.styles
normal = st["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

body = st.add_style("AQS Body", 1)
body.base_style = normal
body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
body.paragraph_format.space_after = Pt(8)
body.paragraph_format.line_spacing = 1.15

for lvl, size in ((1, 18), (2, 14), (3, 12)):
    h = st[f"Heading {lvl}"]
    h.font.color.rgb = BRAND
    h.font.size = Pt(size)
    h.font.name = "Calibri Light"
st["Title"].font.color.rgb = BRAND
st["Title"].font.size = Pt(32)

# Multi-level heading numbering (1 / 1.1 / 1.1.1) linked to the heading styles.
numbering = d.part.numbering_part.element
abstract = parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>'
    + "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading{i+1}"/>'
        f'<w:lvlText w:val="{".".join(f"%{k+1}" for k in range(i+1))}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{432 + 144*i}" w:hanging="{432 + 144*i}"/></w:pPr></w:lvl>'
        for i in range(3)
    )
    + "</w:abstractNum>"
)
numbering.insert(0, abstract)
numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
for lvl in (1, 2, 3):
    ppr = st[f"Heading {lvl}"].element.get_or_add_pPr()
    ppr.insert(0, parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{lvl-1}"/><w:numId w:val="90"/></w:numPr>'))

# ---- cover --------------------------------------------------------------
d.add_paragraph("Architecture Quick Scan", style="Title")
p = d.add_paragraph(style="Subtitle")
p.add_run("Klant: ")
sdt = parse_xml(
    f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Klant"/><w:tag w:val="klant"/><w:showingPlcHdr/></w:sdtPr>'
    f'<w:sdtContent><w:r><w:rPr><w:color w:val="808080"/></w:rPr><w:t>[Klantnaam]</w:t></w:r></w:sdtContent></w:sdt>'
)
p._p.append(sdt)
d.add_paragraph("Versie: {{versie}}", style="AQS Body")
d.add_paragraph("Datum: {{datum}}", style="AQS Body")
d.add_paragraph("Auteur: {{auteur}}", style="AQS Body")

# ---- TOC ------------------------------------------------------------------
d.add_page_break()
toc_head = d.add_paragraph("Inhoudsopgave", style="AQS Body")
toc_head.runs[0].bold = True
toc_head.runs[0].font.size = Pt(16)
toc = d.add_paragraph()
toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'))
toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t>Klik met rechts en kies Veld bijwerken.</w:t></w:r>'))
toc._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))

# New section for the body: own header/footer.
sec = d.add_section(WD_SECTION.NEW_PAGE)
sec.header.is_linked_to_previous = False
sec.footer.is_linked_to_previous = False
hp = sec.header.paragraphs[0]
hp.text = "Architecture Quick Scan | {{klant}}"
hp.runs[0].font.color.rgb = BRAND
fp = sec.footer.paragraphs[0]
fp.add_run("Pagina ")
fp._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
d.sections[0].different_first_page_header_footer = True


def guidance(text):
    p = d.add_paragraph(style="AQS Body")
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


d.add_heading("Inleiding", 1)
d.add_heading("Aanleiding", 2)
guidance("[Beschrijf de aanleiding voor deze quick scan en de vraag van de opdrachtgever.]")
d.add_heading("Doel en scope", 2)
guidance("[Beschrijf doel, scope en afbakening. Wat valt buiten scope?]")
d.add_paragraph("Deze quick scan is uitgevoerd op basis van interviews en documentstudie.", style="AQS Body")
d.add_heading("Huidige situatie", 1)
guidance("[Beschrijf het huidige applicatielandschap, de belangrijkste processen en knelpunten.]")
d.add_paragraph("Belangrijkste knelpunten:", style="AQS Body")
d.add_paragraph("Voorbeeldknelpunt", style="List Bullet")
d.add_heading("Bevindingen", 1)
guidance("[Vat per architectuurthema de bevindingen samen.]")
t = d.add_table(rows=2, cols=4)
t.style = st["Table Grid"]
for i, txt in enumerate(["Thema", "Bevinding", "Risico", "Advies"]):
    c = t.rows[0].cells[i]
    c.text = txt
    c.paragraphs[0].runs[0].bold = True
    c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="004B87"/>'))
for i, txt in enumerate(["Informatiebeveiliging", "...", "Hoog", "..."]):
    t.rows[1].cells[i].text = txt
d.add_heading("Wet- en regelgeving", 1)
guidance("[Benoem relevante wet- en regelgeving (bijv. AVG, BIO, Archiefwet, Woo) en de impact.]")
d.add_heading("Advies", 1)
guidance("[Geef concrete, geprioriteerde adviezen.]")
d.add_heading("Vervolgstappen", 1)
guidance("[Beschrijf de voorgestelde vervolgstappen met planning.]")

d.save(out)

# Localize style IDs like Dutch Word does (names stay English).
import zipfile, io, re
src = open(out, "rb").read()
zin = zipfile.ZipFile(io.BytesIO(src))
buf = io.BytesIO()
zout = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
ids = {"Heading1": "Kop1", "Heading2": "Kop2", "Heading3": "Kop3", "Title": "Titel", "Subtitle": "Ondertitel",
       "ListBullet": "Lijstopsommingsteken", "TableGrid": "Tabelraster", "Normal": "Standaard"}
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename.endswith(".xml"):
        s = data.decode("utf8")
        for a, b in ids.items():
            s = re.sub(rf'(w:(?:styleId|val)="){a}"', rf'\g<1>{b}"', s)
        data = s.encode("utf8")
    zout.writestr(item, data)
zout.close()
open(out, "wb").write(buf.getvalue())
print("wrote", out)
